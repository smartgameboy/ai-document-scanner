"""Format specialists return text units with source locations."""
import base64
import csv
import io
from pathlib import Path

MAX_BYTES = 20 * 1024 * 1024

def decode(payload: str) -> bytes:
    if len(payload) > (MAX_BYTES * 4 // 3 + 8):
        raise ValueError("File exceeds 20 MB limit")
    data = base64.b64decode(payload, validate=True)
    if len(data) > MAX_BYTES:
        raise ValueError("File exceeds 20 MB limit")
    return data

def unit(name, location, text):
    return {"source": Path(name).name, "location": location, "text": text.strip()}

class PDFReaderAgent:
    def read(self, name, data):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are unsupported")
        return [unit(name, f"page {i}", p.extract_text() or "")
                for i, p in enumerate(reader.pages, 1)]

class CSVReaderAgent:
    def read(self, name, data):
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192])
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text), dialect)
        headers = next(rows, [])
        result = []
        for i, row in enumerate(rows, 2):
            if not any(row):
                continue
            pairs = [f"{headers[j] if j < len(headers) else 'column ' + str(j+1)}: {v}"
                     for j, v in enumerate(row)]
            result.append(unit(name, f"record {i} (header is record 1)", " | ".join(pairs)))
        return result

class ExcelReaderAgent:
    def read(self, name, data):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        result = []
        try:
            for sheet in wb:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                for i, row in enumerate(rows, 2):
                    if not any(v is not None for v in row):
                        continue
                    text = " | ".join(f"{headers[j] if j < len(headers) and headers[j] is not None else 'column '+str(j+1)}: {v}"
                                      for j, v in enumerate(row) if v is not None)
                    result.append(unit(name, f"sheet {sheet.title}, row {i}", text))
        finally:
            wb.close()
        return result

class WordReaderAgent:
    def read(self, name, data):
        from docx import Document
        doc = Document(io.BytesIO(data))
        result = [unit(name, f"paragraph {i}", p.text)
                  for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
        for t, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                result.append(unit(name, f"table {t}, row {r}", " | ".join(c.text for c in row.cells)))
        return result
