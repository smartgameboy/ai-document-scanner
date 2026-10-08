import asyncio
import io
import numpy as np
import pytest
from readers import CSVReaderAgent, ExcelReaderAgent, WordReaderAgent, PDFReaderAgent, decode
from manager import ManagerAgent, RAGIndex, chunk_units


def test_csv_quoted_multiline():
    units = CSVReaderAgent().read('a.csv', b'name,notes\nShreya,"AI, RAG\nproject"\n')
    assert len(units) == 1
    assert 'AI, RAG\nproject' in units[0]['text']


def test_excel_and_word():
    from openpyxl import Workbook
    from docx import Document
    wb = Workbook()
    wb.active.append(['Name', 'Topic'])
    wb.active.append(['Shreya', 'RAG'])
    stream = io.BytesIO()
    wb.save(stream)
    assert 'Topic: RAG' in ExcelReaderAgent().read('a.xlsx', stream.getvalue())[0]['text']
    doc = Document()
    doc.add_paragraph('Agentic AI')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'MCP'
    stream = io.BytesIO()
    doc.save(stream)
    units = WordReaderAgent().read('a.docx', stream.getvalue())
    assert [u['text'] for u in units] == ['Agentic AI', 'MCP']


def test_pdf_blank_and_malformed():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    assert PDFReaderAgent().read('a.pdf', stream.getvalue())[0]['text'] == ''
    with pytest.raises(Exception):
        PDFReaderAgent().read('bad.pdf', b'not a pdf')
    with pytest.raises(Exception):
        decode('%%%')


def test_chunk_boundaries():
    text = 'a' * 2401
    units = [{'text': text, 'source': 'a', 'location': 'p1'}]
    chunks = chunk_units(units)
    assert len(chunks) == 3
    assert all(len(c['text']) <= 1200 for c in chunks)
    assert sum(len(c['text']) for c in chunks) - 360 == len(text)


class FakeProvider:
    def embed(self, texts, query=False):
        return np.asarray([[1., 0.] if 'apple' in t else [0., 1.] for t in texts])
    def answer(self, question, evidence):
        assert '[S1]' in evidence
        return 'Apple [S1]'


def test_retrieval_and_citations():
    provider = FakeProvider()
    chunks = [{'text': t, 'source': 'a', 'location': 'row 1'} for t in ['banana', 'apple']]
    index = RAGIndex(chunks, provider)
    answer, hits = index.ask('apple', provider, 1)
    assert hits[0]['text'] == 'apple'
    assert answer == 'Apple [S1]'


def test_real_mcp_roundtrip_and_failure_isolation():
    chunks, trace, errors = asyncio.run(ManagerAgent().ingest([
        ('test.csv', b'name,topic\nShreya,RAG\n'),
        ('bad.pdf', b'bad'), ('other.exe', b'bad')]))
    assert len(chunks) == 1
    assert 'Shreya' in chunks[0]['text']
    assert trace[0]['agent'] == 'read_csv'
    assert len(errors) == 2
